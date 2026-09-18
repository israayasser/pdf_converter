from pathlib import Path
import re
import tempfile
import zipfile

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


TRANSACTION_HEADER = [
    "GMT DDMM HH:MI",
    "TRAN DATE DDMM HH:MI",
    "PAN",
    "DEST",
    "MSG TYPE",
    "FUNC CODE",
    "ACT CODE",
    "TRAN TYPE",
    "TERMINAL ID",
    "RRN",
    "TRAN AMT",
    "INT FEE",
    "NET FEE",
]


def clean(value):
    return "" if value is None else re.sub(r"\s+", " ", str(value)).strip()


def normalize(value):
    return clean(value).upper()


def detect_section(text):
    match = re.search(r"TRAN\s+TYPE\s*:\s*(.+)", text, re.IGNORECASE)
    return clean(match.group(1)) if match else None


def is_transaction_header(row):
    if not row:
        return False
    text = " ".join(normalize(cell) for cell in row)
    required = [
        "GMT", "TRAN DATE", "PAN", "DEST", "MSG TYPE", "FUNC CODE",
        "ACT CODE", "TRAN TYPE", "TERMINAL ID", "RRN", "TRAN AMT",
        "INT FEE", "NET FEE",
    ]
    return all(item in text for item in required)


def find_transaction_table(tables):
    for table in tables:
        for row_index, row in enumerate(table):
            if is_transaction_header(row):
                return table, row_index
    return None, None


def extract_rows(table, header_index):
    rows = []
    for row in table[header_index + 1:]:
        if not row:
            continue
        row = [clean(cell) for cell in row]
        if not any(row) or is_transaction_header(row):
            continue
        row = row[:13]
        rows.append(row + [""] * (13 - len(row)))
    return rows


def make_sheet_name(name, used_names):
    name = re.sub(r'[\\/*?:\[\]]', "-", name).strip() or "Section"
    name = name[:31]
    original = name
    counter = 2
    while name in used_names:
        suffix = f" ({counter})"
        name = original[:31 - len(suffix)] + suffix
        counter += 1
    used_names.add(name)
    return name


def format_sheet(ws):
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1565C0")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for column in ws.columns:
        max_length = max(
            (len(str(cell.value)) for cell in column if cell.value is not None),
            default=0,
        )
        letter = get_column_letter(column[0].column)
        ws.column_dimensions[letter].width = min(max(max_length + 2, 10), 30)


def write_excel(sections, output_path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names = set()

    for section_name, rows in sections.items():
        ws = workbook.create_sheet(make_sheet_name(section_name, used_names))
        ws.append(TRANSACTION_HEADER)
        for row in rows:
            ws.append(row)
        format_sheet(ws)

    workbook.save(output_path)


def process_pdf(pdf_path, progress_callback=None):
    pdf_path = Path(pdf_path)
    if not pdf_path.exists():
        raise FileNotFoundError(f"PDF not found:\n{pdf_path}")

    sections = {}
    current_section = None
    pages_processed = 0

    with pdfplumber.open(pdf_path) as pdf:
        total_pages = len(pdf.pages)

        for page_number, page in enumerate(pdf.pages, start=1):
            pages_processed = page_number

            if progress_callback:
                progress_callback(page_number, total_pages, "Reading PDF...")

            page_text = page.extract_text() or ""
            detected_section = detect_section(page_text)

            if detected_section:
                current_section = detected_section
                sections.setdefault(current_section, [])

            if progress_callback:
                progress_callback(page_number, total_pages, "Reading tables...")

            tables = page.extract_tables()
            table, header_index = find_transaction_table(tables)

            if table is None or current_section is None:
                continue

            sections[current_section].extend(
                extract_rows(table, header_index)
            )

    if not sections:
        raise ValueError("No transaction sections were found.")

    total_rows = sum(len(rows) for rows in sections.values())
    purchase_rows = sum(
        len(rows)
        for section_name, rows in sections.items()
        if "PURCHASE" in section_name.upper()
    )
    expected_purchase_rows = 3866

    return {
        "sections_data": sections,
        "pages": pages_processed,
        "sections": len(sections),
        "rows": total_rows,
        "purchase_rows": purchase_rows,
        "expected_purchase_rows": expected_purchase_rows,
        "purchase_valid": purchase_rows == expected_purchase_rows,
    }


def convert_pdf_to_excel(pdf_path, output_path, progress_callback=None):
    data = process_pdf(pdf_path, progress_callback)
    if progress_callback:
        progress_callback(data["pages"], data["pages"], "Creating Excel file...")

    write_excel(data["sections_data"], output_path)

    return {
        **{key: value for key, value in data.items() if key != "sections_data"},
        "output": str(output_path),
    }


def merge_sections(section_sets):
    merged = {}
    for sections in section_sets:
        for section_name, rows in sections.items():
            merged.setdefault(section_name, []).extend(rows)
    return merged


def process_zip(zip_path, output_mode="merge_all", group_size=2, progress_callback=None):
    zip_path = Path(zip_path)
    if not zip_path.exists():
        raise FileNotFoundError(f"ZIP not found:\n{zip_path}")

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_dir = Path(temp_dir)

        with zipfile.ZipFile(zip_path) as archive:
            pdf_names = sorted(
                name for name in archive.namelist()
                if name.lower().endswith(".pdf") and not name.endswith("/")
            )
            if not pdf_names:
                raise ValueError("No PDF files were found in the ZIP file.")
            archive.extractall(temp_dir)

        results = []
        total_pages = 0
        completed_pages = 0

        for index, name in enumerate(pdf_names, start=1):
            pdf_path = temp_dir / name
            try:
                data = process_pdf(
                    pdf_path,
                    lambda page, total, status, i=index, n=name: progress_callback(
                        completed_pages + page,
                        completed_pages + total,
                        f"PDF {i}/{len(pdf_names)} — {n} — {status}"
                    ) if progress_callback else None,
                )
                results.append((name, data))
                total_pages += data["pages"]
                completed_pages += data["pages"]
            except Exception as exc:
                results.append((name, {"error": str(exc)}))

        valid_results = [item for item in results if "error" not in item[1]]
        if not valid_results:
            errors = "\n".join(f"{name}: {data['error']}" for name, data in results)
            raise ValueError(f"No PDF files could be processed.\n\n{errors}")

        output_dir = zip_path.parent
        stem = zip_path.stem
        outputs = []

        if output_mode == "separate":
            for name, data in valid_results:
                output = output_dir / f"{Path(name).stem}.xlsx"
                write_excel(data["sections_data"], output)
                outputs.append(str(output))

        else:
            groups = (
                [valid_results]
                if output_mode == "merge_all"
                else [
                    valid_results[i:i + group_size]
                    for i in range(0, len(valid_results), group_size)
                ]
            )

            for index, group in enumerate(groups, start=1):
                merged = merge_sections([data["sections_data"] for _, data in group])

                if output_mode == "merge_all":
                    output = output_dir / f"{stem}_merged.xlsx"
                else:
                    first = Path(group[0][0]).stem
                    last = Path(group[-1][0]).stem
                    output = output_dir / f"{stem}_group_{index}_{first}_to_{last}.xlsx"

                write_excel(merged, output)
                outputs.append(str(output))

        warnings = []
        passed = 0
        total_rows = 0

        for name, data in results:
            if "error" in data:
                warnings.append(f"{Path(name).name}: {data['error']}")
                continue

            total_rows += data["rows"]
            if data["purchase_valid"]:
                passed += 1
            else:
                warnings.append(
                    f"{Path(name).name}: Purchase rows "
                    f"{data['purchase_rows']:,} / {data['expected_purchase_rows']:,}"
                )

        return {
            "files": len(pdf_names),
            "passed": passed,
            "warnings": len(warnings),
            "warning_messages": warnings,
            "rows": total_rows,
            "outputs": outputs,
        }


# Keep compatibility with app.py
extract_pdf_to_excel = convert_pdf_to_excel
